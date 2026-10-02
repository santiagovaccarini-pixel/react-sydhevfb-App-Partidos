"""
Genera las filas de Valor Referencial (hoja VR) para cada jugador x categoria de tiempo a partir
de la hoja 'Data GPS Partido' (columna T = jugador, columna DG = categoria de tiempo).

Los niveles se calculan con la logica de la 'Plantilla VR' (hojas 1.3 Proceso_Absolutos y
2.3 Proceso_Relativos), metrica por metrica:
  * Datos raros: si el 10% o menos de los valores cae fuera de 1,5 rangos intercuartilicos (RIC)
    se quitan los que estan fuera de 1,5 RIC; si cae mas del 10%, solo los que estan fuera de 3 RIC.
    Absolutos y relativos por minuto con promedio de 1 o menos no se limpian.
  * Bueno = promedio sin datos raros (absolutos, relativos vs equipo, caidas) o cociente de sumas
    (relativos por minuto).
  * Muy Bueno / Regular = Bueno +/- m*desvio, con m (0,25..1,75) elegido para dejar ~34% de casos
    entre Bueno y ese nivel; Excelente = Bueno + (m + 0,25*i)*desvio, hasta 2 desvios, buscando
    2,5% de casos por encima y 13,5% entre Muy Bueno y Excelente; Malo = Regular - 0,25*desvio.
  * Niveles negativos en metricas que no pueden serlo se llevan a 0 (las caidas pueden ser negativas).
  * Con menos de 5 casos no se arma VR. Los multiplicadores se eligen para que la muestra se reparta como
    una Gauss: Excelente 2,5%, Muy Bueno 13,5%, Bueno 34%, Regular 34%, Malo 16% (Malo = debajo de Regular).
  * Si la muestra se pasa de los estandares (Excelente > 10%, Malo > 20%, niveles negativos o desvio enorme)
    el desvio se reparte a mano con cortes entre valores reales de la muestra (celdas celestes).
  * Solo las caidas pueden tener niveles negativos; se marcan con letra roja.

Ademas junta categorias por jugador segun las celdas PINTADAS en 'Tiempos por jugador' (contiguas
y del mismo color = un grupo; un bloque por grupo con la categoria de mas casos, en amarillo y con
nota en la cantidad de casos). Las celdas con datos raros quitados quedan en naranja.

Uso:  python excel/generar_vr_jugadores.py  <entrada.xlsm>  <salida.xlsm>  [opciones]
        --sin-atipicos         no quitar datos raros
        --permitir-negativos   no llevar a 0 los niveles negativos
        --sin-a-mano           no aplicar el reparto a mano
        --sin-juntar           ignorar las celdas pintadas de 'Tiempos por jugador'
Ver:  excel/VALOR_REFERENCIAL.md
"""
import sys, os, re, zipfile, datetime, math, html, warnings, argparse
warnings.filterwarnings('ignore')
from collections import Counter, defaultdict
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L, column_index_from_string as C

ap = argparse.ArgumentParser()
ap.add_argument('entrada', nargs='?', default='base.xlsm')
ap.add_argument('salida', nargs='?', default='salida.xlsm')
ap.add_argument('--sin-atipicos', action='store_true')
ap.add_argument('--permitir-negativos', action='store_true')
ap.add_argument('--sin-a-mano', action='store_true')
ap.add_argument('--sin-juntar', action='store_true')
ARGS = ap.parse_args()
SRC, DST = ARGS.entrada, ARGS.salida
REVIEW = DST.rsplit('.', 1)[0] + '_resumen.xlsx'

DATA_SHEET, VR_SHEET, PLAYERS_SHEET = 'Data GPS Partido', 'VR', 'Tiempos por jugador'
HDR_ROW, FIRST_DATA_ROW = 13, 15
FIRST_COL, LAST_COL = C('P'), C('DA')          # columnas con formulas Excelente..Malo
CATS = ['Sólo PT', 'Sólo ST', '>=30 y Final ST', '>=10 y <30 ST', '>=85', '>=70 y <85', 'PT + <25']
ITEM_BY_CAT = {'Sólo PT': 'Jugador PT', 'Sólo ST': 'Jugador ST'}   # resto -> Jugador Total
NIVELES = [('Excelente', 5), ('Muy Bueno', 6), ('Bueno', 7), ('Regular', 8), ('Malo', 9), ('Desv. Estándar', None)]
VR_FIRST_FREE_ROW = 166
TODAY_SERIAL = (datetime.date.today() - datetime.date(1899, 12, 30)).days
FILL_YELLOW_RGB = 'FFFFFF00'    # casos juntados de otra categoria (convencion del libro)
FILL_ORANGE_RGB = 'FFF4B183'    # valores atipicos excluidos del calculo
FILL_BLUE_RGB = 'FFBDD7EE'      # metrica con el desvio repartido a mano
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def norm(h):
    return re.sub(r'\s+', ' ', str(h)).strip().casefold() if h is not None else None


def to_num(v):
    if v is None or v == '':
        return np.nan
    if isinstance(v, datetime.timedelta):
        return v.total_seconds() / 86400.0
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6) / 86400.0
    if isinstance(v, datetime.datetime):
        return (v - datetime.datetime(1899, 12, 30)).total_seconds() / 86400.0
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return np.nan


def serial_to_date(v):
    try:
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))).date()
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- 0. partes XML del libro
zin = zipfile.ZipFile(SRC)
wbxml = zin.read('xl/workbook.xml').decode('utf-8')
rels = zin.read('xl/_rels/workbook.xml.rels').decode('utf-8')


def sheet_part(name):
    rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(name), wbxml).group(1)
    m = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or \
        re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
    t = m.group(1)
    return t.lstrip('/') if t.startswith('/') else 'xl/' + t


styles_xml = zin.read('xl/styles.xml').decode('utf-8')
fills_block = re.search(r'<fills count="(\d+)">(.*?)</fills>', styles_xml, re.S)
fills_list = re.findall(r'<fill>.*?</fill>|<fill/>', fills_block.group(2), re.S)
xfs_block = re.search(r'<cellXfs count="(\d+)">(.*?)</cellXfs>', styles_xml, re.S)
xfs_list = re.findall(r'<xf [^>]*?(?:/>|>.*?</xf>)', xfs_block.group(2), re.S)
xf_fill = [int(re.search(r'fillId="(\d+)"', x).group(1)) for x in xfs_list]

# ---------------------------------------------------------------- 1. leer datos
print('Leyendo', SRC)
wbv = load_workbook(SRC, read_only=True, data_only=True)
wbf = load_workbook(SRC, read_only=True, data_only=False)
ws_v, ws_f = wbv[DATA_SHEET], wbf[DATA_SHEET]

hdr13 = next(ws_v.iter_rows(min_row=HDR_ROW, max_row=HDR_ROW, values_only=True))
ncols = len(hdr13)
players = [str(r[0]).strip() for r in wbv[PLAYERS_SHEET].iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
N4 = next(ws_v.iter_rows(min_row=4, max_row=4, min_col=C('N'), max_col=C('N'), values_only=True))[0]

raw = []
for r in ws_v.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
    if r[C('T') - 1] is None:
        continue
    raw.append(r)
print('Filas de datos:', len(raw))
cols_letters = [L(i + 1) for i in range(ncols)]
num = pd.DataFrame([[to_num(v) for v in r] for r in raw], columns=cols_letters)
T = pd.Series([str(r[C('T') - 1]).strip() for r in raw])
DG = pd.Series([r[C('DG') - 1] for r in raw])
U = pd.Series([r[C('U') - 1] for r in raw])
META = pd.DataFrame({'Caso': [r[C('C') - 1] for r in raw], 'Item': [r[C('D') - 1] for r in raw],
                     'Fecha': [r[C('E') - 1] for r in raw], 'Rival': [r[C('N') - 1] for r in raw]})

# ---------------------------------------------------------------- 2. grupos a juntar (celdas pintadas)
merge_groups = {}   # jugador -> {categoria: [categorias del grupo]}
if not ARGS.sin_juntar:
    pl_xml = zin.read(sheet_part(PLAYERS_SHEET)).decode('utf-8')
    ss = zin.read('xl/sharedStrings.xml').decode('utf-8') if 'xl/sharedStrings.xml' in zin.namelist() else ''
    sst = [html.unescape(''.join(re.findall(r'<t[^>]*>(.*?)</t>', si, re.S)))
           for si in re.findall(r'<si>.*?</si>', ss, re.S)]

    def cell_text(c):
        v = re.search(r'<v>(.*?)</v>', c, re.S)
        if v is None:
            t = re.search(r'<is>.*?<t[^>]*>(.*?)</t>', c, re.S)
            return html.unescape(t.group(1)) if t else None
        if 't="s"' in c:
            return sst[int(v.group(1))]
        return html.unescape(v.group(1))

    cats_by_col = {}
    for c in re.findall(r'<c r="[A-Z]+1"[^>]*?(?:/>|>.*?</c>)', pl_xml, re.S):
        col = re.match(r'<c r="([A-Z]+)', c).group(1)
        txt = cell_text(c)
        if txt in CATS:
            cats_by_col[col] = txt
    for row in re.findall(r'<row r="(\d+)"[^>]*>(.*?)</row>', pl_xml, re.S):
        rnum, inner = int(row[0]), row[1]
        if rnum < 2:
            continue
        cells = {re.match(r'<c r="([A-Z]+)', c).group(1): c
                 for c in re.findall(r'<c r="[A-Z]+\d+"[^>]*?(?:/>|>.*?</c>)', inner, re.S)}
        name = cell_text(cells['A']).strip() if 'A' in cells and cell_text(cells['A']) else None
        if not name:
            continue
        # corridas de celdas pintadas contiguas (mismo color) en el orden de las columnas = un grupo
        runs, cur = [], None
        for col in sorted(cats_by_col, key=C):
            cat = cats_by_col[col]
            c = cells.get(col)
            sm = re.search(r' s="(\d+)"', c) if c is not None else None
            fid = xf_fill[int(sm.group(1))] if sm else 0
            if fid not in (0, 1):
                if cur is not None and cur[0] == fid:
                    cur[1].append(cat)
                else:
                    if cur is not None:
                        runs.append(cur)
                    cur = [fid, [cat]]
            elif cur is not None:
                runs.append(cur); cur = None
        if cur is not None:
            runs.append(cur)
        groups = {}
        for fid, cats in runs:
            if len(cats) >= 2:
                for cat in cats:
                    groups[cat] = cats
            else:
                print(f'  Aviso: {name} / {cats[0]} esta pintada sola (sin una categoria vecina del mismo color); se ignora.')
        if groups:
            merge_groups[name] = groups
    print('Jugadores con categorias a juntar:', len(merge_groups))
    for p, g in merge_groups.items():
        seen = []
        for cats in g.values():
            if cats not in seen:
                seen.append(cats); print(f'   {p}: ' + ' + '.join(cats))

# ---------------------------------------------------------------- 3. formulas fila 5..9
frows = {i: next(ws_f.iter_rows(min_row=i, max_row=i, values_only=True)) for i in (5, 6, 7, 8, 9)}


def ftext(c):
    return getattr(c, 'text', c)


def split_args(s):
    args, depth, cur, q = [], 0, '', False
    for ch in s:
        if ch == '"':
            q = not q
        if not q:
            if ch == '(':
                depth += 1
            elif ch == ')':
                depth -= 1
            elif ch == ',' and depth == 0:
                args.append(cur); cur = ''; continue
        cur += ch
    args.append(cur)
    return args


def translate(f):
    """formula Excel (subconjunto usado en filas 5..9) -> expresion Python"""
    f = f.lstrip('=')
    while True:
        m = re.search(r'IF\(', f)
        if not m:
            break
        start = m.end(); depth = 1; i = start
        while depth:
            if f[i] == '(':
                depth += 1
            elif f[i] == ')':
                depth -= 1
            i += 1
        c, a, b = split_args(f[start:i - 1])
        f = f[:m.start()] + f'(({a}) if ({c}) else ({b}))' + f[i:]
    f = re.sub(r'SUBTOTAL\((\d+),\$?([A-Z]+)\$?\d+:\$?[A-Z]+\$?\d+\)', r"ST(\1,'\2')", f)
    f = re.sub(r'\$?([A-Z]+)\$7\b', r"R7('\1')", f)
    f = f.replace('$N$4', 'N4').replace('="%"', '=="%"')
    f = re.sub(r'(?<![=<>!])=(?!=)', '==', f)
    return f


formulas = {}
for j in range(FIRST_COL - 1, LAST_COL):
    col = L(j + 1)
    f7 = ftext(frows[7][j])
    if not (isinstance(f7, str) and f7.startswith('=')):
        continue
    formulas[col] = {i: translate(ftext(frows[i][j])) for i in (5, 6, 7, 8, 9)}
print('Columnas con formula:', len(formulas))


class Calc:
    """Evalua las formulas de las filas 5..9 sobre un subconjunto de filas.
    excl: {columna objetivo: mascara booleana de filas a excluir (atipicos de esa columna)}"""

    def __init__(self, mask, excl=None):
        self.base = np.asarray(mask, dtype=bool)
        self.excl = excl or {}
        self.cur = self.base
        self.cache = {}

    def _use(self, col):
        e = self.excl.get(col)
        self.cur = self.base & ~e if e is not None else self.base

    def ST(self, n, col):
        s = num[col][self.cur].dropna()
        if n == 1:
            return s.mean() if len(s) else float('nan')
        if n == 2:
            return float(len(s))
        if n == 7:
            return s.std(ddof=1) if len(s) > 1 else float('nan')
        if n == 9:
            return s.sum() if len(s) else 0.0
        raise ValueError(n)

    def R7(self, col):
        # referencia a la fila 7 de la misma columna objetivo: misma mascara
        return self._eval(7, col)

    def _eval(self, row, col):
        key = (row, col)
        if key not in self.cache:
            try:
                v = eval(formulas[col][row], {'ST': self.ST, 'R7': self.R7, 'N4': N4})
                if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
                    v = None
            except (ZeroDivisionError, KeyError, TypeError):
                v = None
            self.cache[key] = v
        return self.cache[key]

    def value(self, row, col):
        self._use(col)
        return self._eval(row, col)

    def std(self, col):
        self._use(col)
        v = self.ST(7, col)
        return None if v is None or np.isnan(v) else float(v)

    def count(self, col):
        self.cur = self.base          # casos del grupo, sin descontar atipicos (como SUBTOTAL(2) en AF1)
        return int(self.ST(2, col))


# ---------------------------------------------------------------- 4. validar contra valores cacheados (sin filtro)
cached = {i: next(ws_v.iter_rows(min_row=i, max_row=i, values_only=True)) for i in (5, 6, 7, 8, 9)}
calc_all = Calc(np.ones(len(num), dtype=bool))
bad = []
for col in formulas:
    for i in (5, 6, 7, 8, 9):
        exp = to_num(cached[i][C(col) - 1]); got = calc_all.value(i, col)
        if got is None and np.isnan(exp):
            continue
        if got is None or np.isnan(exp) or abs(got - exp) > 1e-6 * max(1, abs(exp)):
            bad.append((col, i, exp, got))
bad_cols = sorted({b[0] for b in bad}, key=C)
print('Validacion vs Excel (sin filtro): columnas con diferencias =', bad_cols,
      '(P:S usan rangos hasta la fila 5216 en el libro; no van a VR)' if bad_cols else '')

# ---------------------------------------------------------------- 5. mapa de columnas Data -> VR
ws_vr = wbv[VR_SHEET]
vr_hdr = next(ws_vr.iter_rows(min_row=2, max_row=2, values_only=True))
vr_by_norm = {}
for j, h in enumerate(vr_hdr):
    if h is not None and norm(h) not in vr_by_norm:
        vr_by_norm[norm(h)] = L(j + 1)
col_map = {}
unmatched = []
for j in range(C('AE') - 1, LAST_COL):
    col = L(j + 1); h = hdr13[j]
    if h is None or (col not in formulas and col != 'AE'):
        continue
    key = norm(h)
    if key in vr_by_norm:
        col_map[col] = vr_by_norm[key]
    else:
        unmatched.append((col, h))
if any(c in bad_cols for c in col_map):
    raise SystemExit('Hay columnas que van a VR cuya formula no se pudo replicar: ' + str([c for c in col_map if c in bad_cols]))
print('Columnas mapeadas a VR:', len(col_map), '| sin correspondencia en VR:', unmatched)
metric_name = {d: str(hdr13[C(d) - 1]).strip() for d in col_map}


# ---------------------------------------------------------------- 5b. logica de la Plantilla VR
M_STEPS = [0.25, 0.5, 0.75, 1.0, 1.25, 1.5, 1.75]
M_OPTS = {m: 7 - k for k, m in enumerate(M_STEPS)}       # columna AD de la planilla: 7..1 opciones
T_HALF, T_EXC, T_MB = 0.34, 0.025, 0.135                    # X37, X35, X36 de la planilla
MIN_CASOS = 5            # con menos casos no se arma VR
# proporciones de una distribucion normal (Gauss) para cada rango, como en la planilla (X35..X39)
GAUSS = {'Excelente': 0.025, 'Muy Bueno': 0.135, 'Bueno': 0.34, 'Regular': 0.34, 'Malo': 0.16}
GAUSS_MULT = {'Muy Bueno': 1.0, 'Excelente': 2.0, 'Regular': 1.0}   # multiplicadores teoricos (desempate)
M_MAX = float(os.environ.get('VR_M_MAX', '2'))   # hasta cuantos desvios puede ir Excelente / Malo
# estandares: si la muestra se pasa, el reparto del desvio se hace "a mano" con los datos de la muestra
MAX_EXCELENTE = 0.10     # Excelente con mas del 10% de la muestra
MAX_MALO = 0.20          # Malo (debajo de Regular) con mas del 20% de la muestra
MAX_DESVIO_RATIO = 1.5   # desvio de la muestra mas de 1,5 veces el desvio del centro de la muestra ((P84 - P16) / 2)
KIND_NAME = {'abs': 'Absoluto', 'rel': 'Relativo por minuto', 'vseq': 'Relativo vs equipo', 'caida': 'Caída'}
NONNEG = {'abs', 'rel', 'vseq'}
NUM_OF = {L(c): L(c - (C('AU') - C('AF'))) for c in range(C('AU'), C('BH') + 1)}   # AU = AF / DB ...
BASE_ABS = dict(NUM_OF)                                                           # relativo -> su absoluto
BASE_ABS.update({L(c): L(c - (C('BY') - C('AF'))) for c in range(C('BY'), C('CL') + 1)})   # vs equipo -> su absoluto


def metric_kind(col):
    c = C(col)
    if C('AU') <= c <= C('BH'):
        return 'rel'
    if C('BJ') <= c <= C('BW') or C('CN') <= c <= C('DA'):
        return 'caida'
    if C('BY') <= c <= C('CL'):
        return 'vseq'
    return 'abs'


def xl_rank(vals, i):
    """RANK de Excel (orden descendente): 1 + cantidad de valores mayores."""
    return 1 + sum(1 for v in vals if v > vals[i])


def pick_m(kept, b, sd, f, side):
    """Multiplicador de Muy Bueno (side='up') o Regular (side='down'), filas 35..49 de la planilla."""
    n = len(kept)
    Z, Y = [], []
    for m in M_STEPS:
        if side == 'up':
            prop = np.sum((kept >= b) & (kept < b + m * sd)) / n
        else:
            prop = np.sum((kept < b) & (kept >= b - m * sd)) / n
        Z.append(abs(T_HALF - prop)); Y.append(m - f)
    AA = [xl_rank(Z, i) for i in range(len(M_STEPS))]
    AB = [AA[i] + xl_rank(Y, i) if AA[i] == max(AA) else None for i in range(len(M_STEPS))]
    best = max(v for v in AB if v is not None)
    return M_STEPS[AB.index(best)]


def pick_exc(kept, b, sd, m_up):
    """Multiplicador de Excelente: opciones m+0.25, m+0.5, ... hasta 2 (filas 1..3 y 14..15 de la planilla)."""
    n = len(kept); best = None
    for i in range(1, 8):
        mult = m_up + 0.25 * i
        if M_OPTS[m_up] < i or mult > 2:
            break
        exc, mb = b + mult * sd, b + m_up * sd
        err = abs(T_EXC - np.sum(kept >= exc) / n) + abs(T_MB - np.sum((kept >= mb) & (kept < exc)) / n)
        if best is None or err < best[0]:
            best = (err, mult)
    return best[1]


def fit_gauss(kept, b, sd):
    """Elige los multiplicadores (de a 0,25 desvios) para que la muestra se reparta en los rangos lo mas
    parecido posible a una Gauss: Excelente 2,5%, Muy Bueno 13,5%, Bueno 34%, Regular 34%, Malo 16%.
    Rangos como pinta el libro: Excelente >= E; Muy Bueno [MB, E); Bueno [B, MB); Regular [R, B); Malo < R.
    Empates: lo mas cerca de los multiplicadores teoricos (1 y 2 desvios)."""
    grid = [round(0.25 * k, 2) for k in range(1, int(M_MAX / 0.25) + 1)]
    best_up = None
    for mu in grid:
        if mu > 1.75:
            break
        for me in grid:
            if me <= mu:
                continue
            pe = np.mean(kept >= b + me * sd); pmb = np.mean((kept >= b + mu * sd) & (kept < b + me * sd))
            pb = np.mean((kept >= b) & (kept < b + mu * sd))
            err = abs(pe - GAUSS['Excelente']) + abs(pmb - GAUSS['Muy Bueno']) + abs(pb - GAUSS['Bueno'])
            key = (round(err, 12), abs(mu - GAUSS_MULT['Muy Bueno']) + abs(me - GAUSS_MULT['Excelente']), me, mu)
            if best_up is None or key < best_up[0]:
                best_up = (key, mu, me)
    best_lo = None
    for ml in grid:
        if ml + 0.25 > M_MAX + 1e-9:
            break
        pr = np.mean((kept >= b - ml * sd) & (kept < b)); pm = np.mean(kept < b - ml * sd)
        err = abs(pr - GAUSS['Regular']) + abs(pm - GAUSS['Malo'])
        key = (round(err, 12), abs(ml - GAUSS_MULT['Regular']), ml)
        if best_lo is None or key < best_lo[0]:
            best_lo = (key, ml)
    return best_up[1], best_up[2], best_lo[1]


def manual_levels(kept, b):
    """Reparto del desvio "a mano": los cortes se ponen entre valores reales de la muestra (nunca fuera de su
    rango, asi no hay negativos en metricas que no pueden serlo ni un desvio inflado), buscando el reparto de
    Gauss (2,5 / 13,5 / 34 / 34 / 16 %) sin pasar del 10% en Excelente ni del 20% en Malo.
    Devuelve (niveles, desvio_manual)."""
    xs = np.sort(kept); n = len(xs)
    cuts = [k for k in range(n + 1) if k in (0, n) or xs[k - 1] < xs[k]]      # k = valores que quedan debajo del corte
    spread = (np.percentile(xs, 84) - np.percentile(xs, 16)) / 2 or (xs[-1] - xs[0]) / 4
    def thr(k):
        if k == 0:
            # debajo del menor valor, a media distancia del siguiente (igual que los otros cortes), para que
            # ningun redondeo de Excel deje al menor valor en Malo
            nxt = xs[xs > xs[0]]
            return xs[0] - (0.5 * (nxt[0] - xs[0]) if len(nxt) else 0.25 * spread)
        if k == n:
            return xs[-1] + 0.25 * spread
        return (xs[k - 1] + xs[k]) / 2
    below_b = int(np.sum(xs < b))
    # Regular: k valores quedan en Malo
    best_r = None
    for k in cuts:
        if k >= n or k / n > MAX_MALO + 1e-9:
            break
        t = thr(k)
        if t >= b:
            break
        pm, pr = k / n, (below_b - k) / n
        key = (round(abs(pm - GAUSS['Malo']) + abs(pr - GAUSS['Regular']), 12), abs(pm - GAUSS['Malo']))
        if best_r is None or key < best_r[0]:
            best_r = (key, t)
    R = best_r[1] if best_r else min(xs[0], b)
    # Muy Bueno y Excelente
    best_u = None
    for k_mb in cuts:
        t_mb = thr(k_mb)
        if t_mb <= b:
            continue
        for k_e in cuts:
            if k_e < k_mb or (k_e == k_mb and k_e < n):
                continue
            j = n - k_e
            if j / n > MAX_EXCELENTE + 1e-9:
                continue
            pe, pmb, pb = j / n, (k_e - k_mb) / n, (k_mb - below_b) / n
            err = abs(pe - GAUSS['Excelente']) + abs(pmb - GAUSS['Muy Bueno']) + abs(pb - GAUSS['Bueno'])
            key = (round(err, 12), abs(pe - GAUSS['Excelente']), abs(pmb - GAUSS['Muy Bueno']))
            if best_u is None or key < best_u[0]:
                best_u = (key, k_mb, k_e)
    if best_u is None:
        MB = thr(n); E = MB + 0.25 * spread
    else:
        MB = thr(best_u[1])
        E = thr(best_u[2]) if best_u[2] < n else max(b + 2 * (MB - b), thr(n))
    sd_m = spread          # desvio del centro de la muestra: (P84 - P16) / 2
    return {'Excelente': E, 'Muy Bueno': MB, 'Regular': R, 'Malo': R - 0.25 * sd_m}, sd_m


def shares(x, b, lv):
    E, MB, R = lv['Excelente'], lv['Muy Bueno'], lv['Regular']
    return {'Excelente': float(np.mean(x >= E)), 'Muy Bueno': float(np.mean((x >= MB) & (x < E))),
            'Bueno': float(np.mean((x >= b) & (x < MB))), 'Regular': float(np.mean((x >= R) & (x < b))),
            'Malo': float(np.mean(x < R))}


def xl_skew(x):
    n = len(x)
    if n < 3:
        return None
    s = np.std(x, ddof=1)
    if not s:
        return None
    return float(n / ((n - 1) * (n - 2)) * np.sum(((x - np.mean(x)) / s) ** 3))


# ---------------------------------------------------------------- 6. calcular por jugador x categoria
out_rows, review, skipped, outliers, combos = [], [], [], [], []
proceso, sin_vr, revisar = [], [], []
for p in players:
    done = set()
    for cat0 in CATS:
        if cat0 in done:
            continue
        group = merge_groups.get(p, {}).get(cat0)
        if group:
            # un solo bloque por grupo, etiquetado con la categoria que mas casos propios tiene
            # (empate: la primera en el orden de las columnas)
            own_counts = {g: int(((T == p) & (DG == g)).sum()) for g in group}
            done.update(group)
            cat = max(group, key=lambda g: (own_counts[g], -group.index(g)))
            mask = ((T == p) & DG.isin(group)).to_numpy()
            n_own = own_counts[cat]
        else:
            cat = cat0
            mask = ((T == p) & (DG == cat)).to_numpy()
            n_own = int(mask.sum())
        n = int(mask.sum())
        if n == 0:
            skipped.append((p, cat)); continue
        cnt = int((~np.isnan(num['AF'].to_numpy()[np.flatnonzero(mask)])).sum())
        if cnt < MIN_CASOS:
            sin_vr.append({'Jugador': p, 'Categoría': cat, 'Casos': cnt,
                           'Juntada con': ' + '.join(g for g in group if g != cat) if group else '',
                           'Motivo': f'menos de {MIN_CASOS} casos'})
            continue
        note = ''
        if group:
            note = ('Categorías juntadas: ' + ' + '.join(f'{g} ({own_counts[g]})' for g in group)
                    + f' = {n} casos. Se muestra como "{cat}" por ser la de más casos.')

        # --- logica de la Plantilla VR, metrica por metrica
        idx = np.flatnonzero(mask)
        excl, n_out_by_col, levels, floored_cols = {}, {}, {}, set()
        manual_cols, neg_cells = set(), set()
        abs_keep = {}      # columna absoluta -> filas (indices globales) que quedaron despues de limpiar
        abs_blank = set()  # columnas absolutas sin VR (Bueno 0 o desvio 0)
        for dcol in col_map:
            kind = metric_kind(dcol)
            vals = num[dcol].to_numpy()[idx]
            ok = ~np.isnan(vals)
            x = vals[ok]
            levels[dcol] = {nv: None for nv, _ in NIVELES}
            if len(x) < MIN_CASOS:
                if len(x):
                    proceso.append({'Jugador': p, 'Categoría': cat, 'Métrica': metric_name[dcol], 'Columna': dcol,
                                    'Tipo': KIND_NAME[kind], 'Casos': int(len(x)), 'Regla': f'sin VR: menos de {MIN_CASOS} casos'})
                continue
            # D7: promedio simple (absolutos) o cociente de sumas (relativos), sobre todos los datos
            if kind == 'rel':
                nume = num[NUM_OF[dcol]].to_numpy()[idx][ok]; deno = num['DB'].to_numpy()[idx][ok]
                sb, sc = np.nansum(nume), np.nansum(deno)
                d7 = (0.01 / sc if sb == 0 else sb / sc) if sc else np.nan
            elif dcol == 'AE':
                d7 = float(np.mean(x)) * 1440          # Tiempo en minutos (en el libro esta en dias)
            else:
                d7 = float(np.mean(x))
            exempt = kind in ('abs', 'rel') and not np.isnan(d7) and d7 <= 1
            keep_local = ok.copy()
            rule, lo, hi, q1, q3, p15 = 'sin limpieza (promedio <= 1)', None, None, None, None, None
            base = BASE_ABS.get(dcol)
            if base is not None and base in abs_keep:
                # relativos (absoluto / minutos) y relativos vs equipo salen del absoluto: usan los mismos partidos
                # que quedaron en el absoluto, asi un partido quitado en el absoluto tambien sale del relativo
                in_abs = np.isin(idx, abs_keep[base])
                keep_local = ok & in_abs
                rule = f'mismos partidos que {metric_name[base]}'
                if kind == 'vseq' and not ARGS.sin_atipicos and keep_local.sum() >= MIN_CASOS:
                    xv = vals[keep_local]
                    q1, q3 = np.percentile(xv, [25, 75]); iqr = q3 - q1
                    tol = 1e-9 * max(abs(q1), abs(q3), iqr, 1e-300)
                    p15 = float(np.sum((xv < q1 - 1.5 * iqr - tol) | (xv > q3 + 1.5 * iqr + tol)) / len(xv))
                    k = 1.5 if p15 <= 0.1 + 1e-12 else 3.0
                    lo, hi = q1 - k * iqr, q3 + k * iqr
                    keep_local = keep_local & (vals >= lo - tol) & (vals <= hi + tol)
                    rule += f' + {k:g} RIC'
            elif ARGS.sin_atipicos:
                rule = 'sin limpieza (--sin-atipicos)'
            elif not exempt:
                q1, q3 = np.percentile(x, [25, 75])       # = CUARTIL / QUARTILE de Excel
                iqr = q3 - q1
                tol = 1e-9 * max(abs(q1), abs(q3), iqr, 1e-300)   # limites inclusivos sin errores de redondeo
                p15 = float(np.sum((x < q1 - 1.5 * iqr - tol) | (x > q3 + 1.5 * iqr + tol)) / len(x))
                k = 1.5 if p15 <= 0.1 + 1e-12 else 3.0
                lo, hi = q1 - k * iqr, q3 + k * iqr
                keep_local = ok & (vals >= lo - tol) & (vals <= hi + tol)
                rule = f'{k:g} RIC' + (' (más del 10% fuera de 1,5 RIC)' if k == 3 else '')
            removed = idx[ok & ~keep_local]
            if kind == 'abs':
                abs_keep[dcol] = idx[keep_local]
            if len(removed):
                m_ = np.zeros(len(num), dtype=bool); m_[removed] = True
                excl[dcol] = m_; n_out_by_col[dcol] = len(removed)
                for ridx in removed:
                    outliers.append({'Jugador': p, 'Categoría VR': cat,
                                     'Categoría del caso': DG[ridx], 'Item': META.Item[ridx],
                                     'Caso': META.Caso[ridx], 'Fecha': META.Fecha[ridx], 'Rival': META.Rival[ridx],
                                     'Métrica': metric_name[dcol], 'Columna': dcol, 'Regla': rule,
                                     'Valor': float(num[dcol][ridx]), 'Límite inferior': lo, 'Límite superior': hi,
                                     'Q1': q1, 'Q3': q3, 'n': int(ok.sum())})
            kept = vals[keep_local]
            # F7 (Bueno)
            if kind == 'rel':
                nume = num[NUM_OF[dcol]].to_numpy()[idx][keep_local]; deno = num['DB'].to_numpy()[idx][keep_local]
                b = float(np.nansum(nume) / np.nansum(deno)) if np.nansum(deno) else None
            else:
                # absolutos, relativos vs equipo y caidas: promedio de los valores que quedan. En las caidas no se usa la
                # formula de la fila 7 (cociente de sumas): en las caidas pp da una fraccion mientras los valores de cada
                # partido estan en puntos porcentuales, y en las caidas relativas puede quedar fuera del rango de los datos.
                b = float(np.mean(kept)) if len(kept) else None
            sd = float(np.std(kept, ddof=1)) if len(kept) > 1 else None
            info = {'Jugador': p, 'Categoría': cat, 'Métrica': metric_name[dcol], 'Columna': dcol, 'Tipo': KIND_NAME[kind],
                    'Casos': int(ok.sum()), '% fuera de 1,5 RIC': p15, 'Regla': rule, 'Valores quitados': int(len(removed)),
                    'Bueno': b, 'Desv. Estándar': sd,
                    'Asimetría': xl_skew(kept), 'F10 (P95 en desvíos)': None, 'F11 (P5 en desvíos)': None,
                    'Mult. Muy Bueno': None, 'Mult. Excelente': None, 'Mult. Regular': None, 'Mult. Malo': None,
                    'Niveles llevados a 0': 0, 'Mult. planilla (MB/Exc/Reg)': None,
                    '% Excelente': None, '% Muy Bueno': None, '% Bueno': None, '% Regular': None, '% Malo': None}
            # como la planilla: si Bueno es 0 o el desvio no se puede calcular (o es 0), el VR de la metrica queda vacio
            computable = b is not None and b != 0 and sd is not None and sd > 0 and len(kept) > 1
            base = BASE_ABS.get(dcol)
            if base is not None and base in abs_keep and base in abs_blank:
                computable = False          # si el absoluto no tiene VR, su relativo tampoco
            if kind == 'abs' and not computable:
                abs_blank.add(dcol)
            if computable:
                levels[dcol]['Bueno'] = b
                levels[dcol]['Desv. Estándar'] = sd
            else:
                info['Regla'] = rule + (' | sin VR: su absoluto no tiene VR' if base in abs_blank and kind != 'abs' else ' | sin VR: Bueno = 0 o desvío 0')
            if computable:
                p95, p5 = np.percentile(kept, 95), np.percentile(kept, 5)
                f10 = (float(np.mean(kept[kept >= p95])) - b) / sd
                f11 = (b - float(np.mean(kept[kept <= p5]))) / sd
                t_up = pick_m(kept, b, sd, f10, 'up')                    # seleccion de la planilla (referencia)
                t_ex = pick_exc(kept, b, sd, t_up)
                t_lo = pick_m(kept, b, sd, f11, 'down')
                m_up, m_ex, m_lo = fit_gauss(kept, b, sd)                   # ajuste a la distribucion de Gauss
                m_ma = m_lo + 0.25
                lv = {'Excelente': b + m_ex * sd, 'Muy Bueno': b + m_up * sd,
                      'Regular': b - m_lo * sd, 'Malo': b - m_ma * sd}
                # estandares de la muestra: si se pasa, reparto del desvio a mano
                p84, p16 = np.percentile(kept, [84, 16]); sig_r = (p84 - p16) / 2
                sh = shares(kept, b, lv)
                motivos = []
                if sh['Excelente'] > MAX_EXCELENTE + 1e-9:
                    motivos.append(f"Excelente {sh['Excelente']:.0%}")
                if sh['Malo'] > MAX_MALO + 1e-9:
                    motivos.append(f"Malo {sh['Malo']:.0%}")
                if kind in NONNEG and min(lv.values()) < 0:
                    motivos.append('niveles negativos')
                if sig_r <= 0 or sd / sig_r > MAX_DESVIO_RATIO:
                    motivos.append('desvío enorme (' + (f'{sd / sig_r:.1f} veces el del centro' if sig_r > 0 else 'centro sin variación') + ')')
                info['Distribución'] = 'Gauss con el desvío de la muestra'
                info['Desvío de la muestra'] = sd
                if motivos and not ARGS.sin_a_mano:
                    lv, sd_m = manual_levels(kept, b)
                    levels[dcol]['Desv. Estándar'] = sd_m
                    info['Distribución'] = 'a mano'; info['Motivo a mano'] = ', '.join(motivos)
                    info['Desv. Estándar'] = sd_m
                    m_up = m_ex = m_lo = m_ma = None
                    manual_cols.add(dcol)
                    revisar.append({'Jugador': p, 'Categoría': cat, 'Métrica': metric_name[dcol], 'Columna': dcol,
                                    'Casos': int(len(kept)), 'Motivo': ', '.join(motivos),
                                    'Desvío de la muestra': sd, 'Desvío a mano': sd_m})
                if kind in NONNEG and not ARGS.permitir_negativos:
                    for k_, v_ in lv.items():
                        if v_ < 0:
                            lv[k_] = 0.0; info['Niveles llevados a 0'] += 1; floored_cols.add(dcol)
                if kind == 'caida':
                    for k_, v_ in lv.items():
                        if v_ < 0:
                            neg_cells.add((dcol, k_))
                    if b < 0:
                        neg_cells.add((dcol, 'Bueno'))
                levels[dcol].update(lv)
                sk, sf = shares(kept, b, lv), shares(x, b, lv)
                info.update({'F10 (P95 en desvíos)': f10, 'F11 (P5 en desvíos)': f11, 'Mult. Muy Bueno': m_up,
                             'Mult. Excelente': m_ex, 'Mult. Regular': m_lo, 'Mult. Malo': m_ma,
                             'Mult. planilla (MB/Exc/Reg)': f'{t_up:g} / {t_ex:g} / {t_lo:g}',
                             **{f'% {k}': v for k, v in sk.items()},
                             **{f'% {k} (con datos raros)': v for k, v in sf.items()}})
            proceso.append(info)
        puesto = Counter(u for u in U[mask] if u).most_common(1)
        puesto = puesto[0][0] if puesto else ''
        item = ITEM_BY_CAT.get(cat, 'Jugador Total')
        juntada = ' + '.join(f'{g} ({own_counts[g]})' for g in group if g != cat) if group else ''
        combos.append({'Jugador': p, 'Categoría': cat, 'Casos propios': n_own, 'Casos usados': n,
                       'Juntada con': juntada, 'Nota en VR': note,
                       'Métricas con atípicos': len(n_out_by_col), 'Valores atípicos excluidos': sum(n_out_by_col.values()),
                       'Métricas con niveles llevados a 0': len(floored_cols), 'Métricas repartidas a mano': len(manual_cols)})
        for nivel, _ in NIVELES:
            row = {'C': item, 'D': p, 'E': puesto, 'F': cat, 'G': nivel, 'H': cnt, 'B': TODAY_SERIAL}
            for dcol, vcol in col_map.items():
                row[vcol] = levels[dcol][nivel]
            row['_key'] = f'{item}{p}{puesto}{cat}{nivel}'
            row['_merged'] = bool(group)
            row['_note'] = note
            row['_outcols'] = {col_map[d] for d in n_out_by_col}
            row['_manualcols'] = {col_map[d] for d in manual_cols}
            row['_negcols'] = {col_map[d] for d, lvl in neg_cells if lvl == nivel}
            out_rows.append(row)
            review.append({'Item': item, 'Nombre': p, 'Puesto': puesto, 'Minutos': cat, 'Categoría': nivel, 'Cuenta': cnt,
                           'Juntada con': juntada, 'Atípicos excluidos': combos[-1]['Valores atípicos excluidos'],
                           **{metric_name[d]: row[v] for d, v in col_map.items()}})
print(f'Combinaciones con VR: {len(combos)}  (filas VR: {len(out_rows)});  sin casos: {len(skipped)};  con menos de {MIN_CASOS} casos (sin VR): {len(sin_vr)}')
_pr = pd.DataFrame(proceso)
if len(_pr):
    _ok = _pr['% Excelente'].notna()
    print('Reparto medio de la muestra por rango (objetivo Gauss 2,5 / 13,5 / 34 / 34 / 16 %): ' +
          ' / '.join(f"{100 * _pr.loc[_ok, f'% {k}'].mean():.1f}" for k in GAUSS) +
          f" | metricas con Excelente > 10%: {int((_pr.loc[_ok, '% Excelente'] > 0.10).sum())}, con Malo > 20%: {int((_pr.loc[_ok, '% Malo'] > 0.20).sum())} de {int(_ok.sum())}")
print(f'Combinaciones juntadas: {sum(1 for c in combos if c["Juntada con"])};  valores atipicos excluidos: {len(outliers)}'
      f' en {sum(1 for c in combos if c["Valores atípicos excluidos"])} combinaciones')
print(f'Niveles llevados a 0 (metricas que no pueden ser negativas): {sum(i.get("Niveles llevados a 0") or 0 for i in proceso)}')
print(f'Metricas con el desvio repartido a mano: {len(revisar)}')

with pd.ExcelWriter(REVIEW) as xw:
    pd.DataFrame(review).to_excel(xw, sheet_name='VR', index=False)
    pd.DataFrame(combos).to_excel(xw, sheet_name='Combinaciones', index=False)
    pd.DataFrame(proceso).to_excel(xw, sheet_name='Proceso', index=False)
    pd.DataFrame(sin_vr, columns=['Jugador', 'Categoría', 'Casos', 'Juntada con', 'Motivo']).to_excel(xw, sheet_name='Sin VR', index=False)
    pd.DataFrame(revisar, columns=['Jugador', 'Categoría', 'Métrica', 'Columna', 'Casos', 'Motivo', 'Desvío de la muestra', 'Desvío a mano']).to_excel(xw, sheet_name='A mano', index=False)
    df_out = pd.DataFrame(outliers)
    if len(df_out):
        df_out['Fecha'] = pd.to_datetime(df_out['Fecha'], errors='coerce').dt.date
    df_out.to_excel(xw, sheet_name='Atípicos', index=False)
    pd.DataFrame({'Leyenda': [
        'Lógica de la Plantilla VR (hojas 1.3 Proceso_Absolutos y 2.3 Proceso_Relativos), métrica por métrica y por jugador y categoría:',
        '1) Datos raros: si el 10% o menos de los valores cae fuera de 1,5 rangos intercuartílicos (RIC), se quitan los que están fuera de 1,5 RIC; si cae más del 10%, se quitan sólo los que están fuera de 3 RIC. Absolutos con promedio de 1 o menos no se limpian. Los relativos por minuto usan los mismos partidos que quedaron en su absoluto (relativo = absoluto / minutos), y los relativos vs equipo también, más su propia limpieza.',
        '2) Bueno = promedio sin datos raros (absolutos, relativos vs equipo y caídas); cociente de sumas de los casos que quedan (relativos por minuto). En las caídas no se usa la fórmula de la fila 7: en las caídas pp da una fracción y los valores de cada partido están en puntos porcentuales.',
        '3) Desv. Estándar = desvío sin datos raros.',
        '4) Muy Bueno, Excelente y Regular = Bueno ± m·desvío, con m de a 0,25 desvíos elegido para que la muestra se reparta en los rangos como una distribución normal (Gauss): Excelente 2,5%, Muy Bueno 13,5%, Bueno 34%, Regular 34%, Malo 16%. Los rangos son los que pinta el libro: Excelente desde Excelente para arriba, Malo debajo de Regular. Excelente puede ir hasta 2 desvíos; en empates se elige lo más cercano a 1 y 2 desvíos.',
        '5) Malo = Regular − 0,25·desvío, como en la planilla (el libro pinta Malo todo lo que está debajo de Regular).',
        f'5b) Reparto a mano: si con el desvío de la muestra Excelente queda con más del {MAX_EXCELENTE:.0%}, Malo con más del {MAX_MALO:.0%}, algún nivel negativo (salvo caídas) o el desvío es más de {MAX_DESVIO_RATIO:g} veces el desvío del centro de la muestra ((P84 − P16) / 2), los cortes se ponen entre valores reales de la muestra buscando el reparto de Gauss sin pasar esos topes. El Desv. Estándar escrito es el del centro de la muestra: (P84 − P16) / 2. Celdas en celeste; detalle en la hoja A mano.',
        '5c) Negativos: sólo las caídas pueden tener niveles negativos; cuando cumplen todo lo anterior se escriben y se marcan con letra roja.',
        '6) La hoja Proceso trae también los multiplicadores que elegía la planilla, como referencia, y el reparto real de la muestra en cada rango, con y sin datos raros.',
        '7) En métricas que no pueden ser negativas (absolutos, relativos por minuto, relativos vs equipo, Tiempo) un nivel que da negativo se lleva a 0. Las caídas pueden ser negativas.',
        f'8) Con menos de {MIN_CASOS} casos no se arma VR: la combinación no se escribe (hoja Sin VR) y una métrica con menos de {MIN_CASOS} valores queda vacía.',

        'Colores en VR: amarillo = categoría calculada juntando casos de otra(s) categoría(s); naranja = métrica a la que se le quitaron datos raros; celeste = métrica con el desvío repartido a mano (tiene prioridad sobre naranja y amarillo); letra roja = nivel negativo en una caída.',
        'Hoja Proceso: detalle por métrica (regla aplicada, multiplicadores elegidos, asimetría). Hoja Atípicos: cada valor quitado.']}).to_excel(xw, sheet_name='Leyenda', index=False)
print('Resumen de control:', REVIEW)

# ---------------------------------------------------------------- 7. estilos: variantes con relleno amarillo / naranja
fills_new = list(fills_list)
def fill_id_for(rgb):
    for i, f in enumerate(fills_new):
        if f'rgb="{rgb}"' in f and 'patternType="solid"' in f:
            return i
    fills_new.append(f'<fill><patternFill patternType="solid"><fgColor rgb="{rgb}"/><bgColor indexed="64"/></patternFill></fill>')
    return len(fills_new) - 1
FILL_Y, FILL_O, FILL_B = fill_id_for(FILL_YELLOW_RGB), fill_id_for(FILL_ORANGE_RGB), fill_id_for(FILL_BLUE_RGB)
fonts_block = re.search(r'<fonts count="(\d+)"([^>]*)>(.*?)</fonts>', styles_xml, re.S)
fonts_list = re.findall(r'<font>.*?</font>|<font/>', fonts_block.group(3), re.S)
fonts_new = list(fonts_list)
red_font_cache = {}
def red_font(fid):
    if fid not in red_font_cache:
        f = fonts_new[fid] if fid < len(fonts_new) else '<font/>'
        if f == '<font/>':
            f = '<font></font>'
        f = re.sub(r'<color [^>]*/>', '', f)
        f = re.sub(r'(<font>(?:<b/>)?(?:<i/>)?(?:<strike/>)?(?:<condense[^>]*/>)?(?:<extend[^>]*/>)?(?:<outline[^>]*/>)?(?:<shadow[^>]*/>)?(?:<u[^>]*/>)?(?:<vertAlign[^>]*/>)?(?:<sz [^>]*/>)?)',
                   r'\1<color rgb="FFC00000"/>', f, count=1)
        fonts_new.append(f); red_font_cache[fid] = len(fonts_new) - 1
    return red_font_cache[fid]
xfs_new = list(xfs_list)
variant_cache = {}
def styled(s_attr, fill=None, red=False):
    """id de estilo igual a s_attr pero con el relleno dado y/o letra roja"""
    key = (s_attr, fill, red)
    if key not in variant_cache:
        base = xfs_new[int(s_attr)] if s_attr else '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        x = base
        if fill is not None:
            x = re.sub(r'fillId="\d+"', f'fillId="{fill}"', x, count=1)
            x = re.sub(r'applyFill="\d"', 'applyFill="1"', x, count=1) if 'applyFill=' in x else x.replace('<xf ', '<xf applyFill="1" ', 1)
        if red:
            fid = int(re.search(r'fontId="(\d+)"', x).group(1))
            x = re.sub(r'fontId="\d+"', f'fontId="{red_font(fid)}"', x, count=1)
            x = re.sub(r'applyFont="\d"', 'applyFont="1"', x, count=1) if 'applyFont=' in x else x.replace('<xf ', '<xf applyFont="1" ', 1)
        xfs_new.append(x); variant_cache[key] = str(len(xfs_new) - 1)
    return variant_cache[key]

# ---------------------------------------------------------------- 8. inyectar en el XML de la hoja VR
vr_part = sheet_part(VR_SHEET)
print('Parte XML de VR:', vr_part)
sx = zin.read(vr_part).decode('utf-8')


def fmt_num(v):
    v = float(v)
    return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)


def render_cell(ref, s_attr, value, existing=None):
    s = f' s="{s_attr}"' if s_attr else ''
    if existing is not None and '<f' in existing:            # celda con formula (columna A): solo actualizar <v>
        fm = re.search(r'<f[^>]*?(?:/>|>.*?</f>)', existing, re.S).group(0)
        return f'<c r="{ref}"{s} t="str">{fm}<v>{html.escape(str(value), quote=False)}</v></c>'
    if value is None:
        return f'<c r="{ref}"{s}/>'
    if isinstance(value, str):
        return f'<c r="{ref}"{s} t="inlineStr"><is><t>{html.escape(value, quote=False)}</t></is></c>'
    return f'<c r="{ref}"{s}><v>{fmt_num(value)}</v></c>'


row_re = re.compile(r'<row r="(\d+)"([^>]*?)(?:/>|>(.*?)</row>)', re.S)
cell_re = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
rows_xml = {int(m.group(1)): m for m in row_re.finditer(sx)}
needed = range(VR_FIRST_FREE_ROW, VR_FIRST_FREE_ROW + len(out_rows))
missing = [r for r in needed if r not in rows_xml]
if missing:
    raise SystemExit(f'La hoja VR no tiene filas pre-formateadas suficientes; faltan {len(missing)} (desde {missing[0]}).')

replacements = []
for r, data in zip(needed, out_rows):
    m = rows_xml[r]
    attrs, inner = m.group(2), m.group(3) or ''
    cells = {cm.group(1): (cm.group(3), cm.group(0)) for cm in cell_re.finditer(inner)}

    row_style = re.search(r' s="(\d+)"', attrs)
    row_style = row_style.group(1) if row_style else None

    def s_of(col):
        a = cells.get(col, ('', ''))[0]
        sm = re.search(r's="(\d+)"', a)
        if sm:
            return sm.group(1)
        # la fila plantilla no tiene esa celda: usar el estilo de la fila (o el de la celda C vecina)
        if col in cells or col == 'C':
            return row_style
        return s_of('C') or row_style

    new_cells = {col: whole for col, (a, whole) in cells.items()}
    a_existing = cells.get('A', (None, None))[1]
    if a_existing and '<f' in a_existing:
        new_cells['A'] = render_cell(f'A{r}', s_of('A'), data['_key'], existing=a_existing)
    else:
        sa = s_of('A'); sattr = f' s="{sa}"' if sa else ''
        new_cells['A'] = f'<c r="A{r}"{sattr} t="str"><f>C{r}&amp;D{r}&amp;E{r}&amp;F{r}&amp;G{r}</f><v>{html.escape(data["_key"], quote=False)}</v></c>'
    for col, v in data.items():
        if col.startswith('_') or col == 'A':
            continue
        s = s_of(col)
        fill = FILL_B if col in data['_manualcols'] else FILL_O if col in data['_outcols'] else FILL_Y if data['_merged'] else None
        red = col in data['_negcols']
        if fill is not None or red:
            s = styled(s, fill, red)
        new_cells[col] = render_cell(f'{col}{r}', s, v)
    ordered = ''.join(new_cells[c] for c in sorted(new_cells, key=C))
    replacements.append((m.start(), m.end(), f'<row r="{r}"{attrs}>{ordered}</row>'))

parts, last = [], len(sx)
for start, end, new in sorted(replacements, reverse=True):
    parts.append(sx[end:last]); parts.append(new); last = start
parts.append(sx[:last])
sx_new = ''.join(reversed(parts))

# estilos nuevos
styles_new = styles_xml
if len(fills_new) != len(fills_list):
    styles_new = styles_new.replace(fills_block.group(0), f'<fills count="{len(fills_new)}">' + ''.join(fills_new) + '</fills>', 1)
if len(fonts_new) != len(fonts_list):
    styles_new = styles_new.replace(fonts_block.group(0), f'<fonts count="{len(fonts_new)}"{fonts_block.group(2)}>' + ''.join(fonts_new) + '</fonts>', 1)
if len(xfs_new) != len(xfs_list):
    styles_new = styles_new.replace(xfs_block.group(0), f'<cellXfs count="{len(xfs_new)}">' + ''.join(xfs_new) + '</cellXfs>', 1)
print(f'Estilos: {len(xfs_new) - len(xfs_list)} variantes de relleno agregadas')

# ---------------------------------------------------------------- 9. notas en la celda de casos (H) de las categorias juntadas
notes = [(f'H{r}', d['_note']) for r, d in zip(needed, out_rows) if d['_note']]
changed_parts = {}
if notes:
    sheet_rels_name = re.sub(r'worksheets/(sheet\d+\.xml)$', r'worksheets/_rels/\1.rels', vr_part)
    srels = zin.read(sheet_rels_name).decode('utf-8') if sheet_rels_name in zin.namelist() else ''
    def rel_target(kind):
        m = re.search(r'<Relationship [^>]*Type="[^"]*/%s"[^>]*Target="([^"]+)"' % kind, srels) or \
            re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Type="[^"]*/%s"' % kind, srels)
        if not m:
            return None
        t = m.group(1)
        return t.lstrip('/') if t.startswith('/') else 'xl/worksheets/' + t if not t.startswith('../') else 'xl/' + t[3:]
    comments_part, vml_part = rel_target('comments'), rel_target('vmlDrawing')
    if not comments_part or not vml_part or comments_part not in zin.namelist() or vml_part not in zin.namelist():
        print('Aviso: la hoja VR no tiene parte de comentarios/VML; las notas de categorias juntadas quedan solo en el resumen.')
    else:
        cx = zin.read(comments_part).decode('utf-8')
        vx = zin.read(vml_part).decode('utf-8')
        existing_refs = set(re.findall(r'<comment ref="([A-Z]+\d+)"', cx))
        author = 'Generador VR'
        authors = re.findall(r'<author>(.*?)</author>', cx, re.S)
        if author in authors:
            aid = authors.index(author)
        else:
            aid = len(authors)
            cx = cx.replace('</authors>', f'<author>{author}</author></authors>', 1)
        sids = [int(x) for x in re.findall(r'id="_x0000_s(\d+)"', vx)]
        next_sid = (max(sids) + 1) if sids else 1025
        zmax = max([int(z) for z in re.findall(r'z-index:(\d+)', vx)] or [0])
        new_comments, new_shapes = [], []
        for ref, text in notes:
            if ref in existing_refs:
                print(f'Aviso: {ref} ya tenia una nota; no se reemplaza.')
                continue
            col_i, row_i = C(re.match(r'[A-Z]+', ref).group(0)), int(re.sub(r'[A-Z]+', '', ref))
            t = html.escape(text, quote=False)
            new_comments.append(
                f'<comment ref="{ref}" authorId="{aid}" shapeId="0"><text>'
                f'<r><rPr><b/><sz val="9"/><color indexed="81"/><rFont val="Tahoma"/><family val="2"/></rPr><t>{author}:</t></r>'
                f'<r><rPr><sz val="9"/><color indexed="81"/><rFont val="Tahoma"/><family val="2"/></rPr><t xml:space="preserve">\n{t}</t></r>'
                f'</text></comment>')
            zmax += 1
            new_shapes.append(
                f'<v:shape id="_x0000_s{next_sid}" type="#_x0000_t202" style=\'position:absolute;margin-left:500pt;margin-top:10pt;'
                f'width:300pt;height:60pt;z-index:{zmax};visibility:hidden\' fillcolor="infoBackground [80]" strokecolor="none [81]" o:insetmode="auto">'
                f'<v:fill color2="infoBackground [80]"/><v:shadow color="none [81]" obscured="t"/><v:path o:connecttype="none"/>'
                f'<v:textbox style=\'mso-direction-alt:auto\'><div style=\'text-align:left\'></div></v:textbox>'
                f'<x:ClientData ObjectType="Note"><x:MoveWithCells/><x:SizeWithCells/>'
                f'<x:Anchor>{col_i}, 15, {max(row_i - 2, 0)}, 10, {col_i + 5}, 15, {row_i + 3}, 4</x:Anchor>'
                f'<x:AutoFill>False</x:AutoFill><x:Row>{row_i - 1}</x:Row><x:Column>{col_i - 1}</x:Column></x:ClientData></v:shape>')
            next_sid += 1
        if new_comments:
            cx = cx.replace('</commentList>', ''.join(new_comments) + '</commentList>', 1)
            vx = vx.replace('</xml>', ''.join(new_shapes) + '</xml>', 1)
            changed_parts[comments_part] = cx
            changed_parts[vml_part] = vx
        print(f'Notas agregadas en VR: {len(new_comments)}')

# recalculo completo al abrir (para que A y las VLOOKUP de 'Data GPS Partido' tomen las filas nuevas)
wbxml_new = wbxml if 'fullCalcOnLoad' in wbxml else re.sub(r'<calcPr([^>]*?)/>', r'<calcPr\1 fullCalcOnLoad="1"/>', wbxml, count=1)
changed_parts.update({vr_part: sx_new, 'xl/workbook.xml': wbxml_new, 'xl/styles.xml': styles_new})

zout = zipfile.ZipFile(DST, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    if item.filename in changed_parts:
        zout.writestr(item, changed_parts[item.filename].encode('utf-8'))
    else:
        zout.writestr(item, zin.read(item.filename))
zout.close()
print('Archivo generado:', DST)
